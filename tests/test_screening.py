"""End-to-end tests for site screening, using small synthetic reference files that go
through the same prepare steps as the real ONS and OS downloads."""

import io
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

from gbeg.screening import prepare, rules
from gbeg.screening.inputs import detect_columns, parse_text_location, read_table
from gbeg.screening.layers import ReferenceData
from gbeg.screening.screen import screen

CFG = rules.load_config()

# A 2 km x 2 km "town" (urban) next to a 2 km x 2 km "countryside" (rural) square.
TOWN = (350_000, 450_000, 352_000, 452_000)
COUNTRY = (352_000, 450_000, 354_000, 452_000)


# ---- input reading ---------------------------------------------------------------------
@pytest.mark.parametrize(
    ("headers", "expected"),
    [
        (["Site Name", "Easting", "Northing", "kVA"],
         {"name": "Site Name", "easting": "Easting", "northing": "Northing", "kva": "kVA"}),
        (["ID", "Lat", "Long", "Noise level dB(A)"],
         {"name": "ID", "lat": "Lat", "lon": "Long", "noise_7m": "Noise level dB(A)"}),
        (["Name", "Post Code", "Grid Ref", "Generator size"],
         {"name": "Name", "postcode": "Post Code", "gridref": "Grid Ref",
          "kva": "Generator size"}),
        (["Site", "X_COORD", "Y_COORD"],
         {"name": "Site", "easting": "X_COORD", "northing": "Y_COORD"}),
    ],
)
def test_detect_columns(headers, expected):
    cols = detect_columns(pd.DataFrame(columns=headers))
    for role, col in expected.items():
        assert cols.get(role) == col


@pytest.mark.parametrize(
    ("text", "method", "easting"),
    [
        ("SJ 838 980", "grid reference", 383_800),
        ("53.4794, -2.2453", "lat/long", 383_818),
        ("-2.2453, 53.4794", "lat/long", 383_818),
        ("383800 398000", "easting/northing", 383_800),
    ],
)
def test_parse_text_location(text, method, easting):
    loc = parse_text_location(text, None)
    assert loc.method == method
    assert loc.easting == pytest.approx(easting, abs=50)


def test_postcode_without_lookup_explains_why():
    loc = parse_text_location("M1 1AE", None)
    assert loc.easting is None
    assert "Postcode lookup not loaded" in loc.issue


def test_read_table_handles_excel_and_semicolons(tmp_path):
    df = pd.DataFrame({"Site": ["A"], "Grid ref": ["SJ 838 980"]})
    buf = io.BytesIO()
    df.to_excel(buf, index=False)
    assert read_table(buf.getvalue(), "sites.xlsx").iloc[0]["Grid ref"] == "SJ 838 980"
    semicolon = b"Site;Grid ref\nA;SJ 838 980\n"
    assert read_table(semicolon, "sites.csv").columns.tolist() == ["Site", "Grid ref"]


# ---- rules -----------------------------------------------------------------------------
@pytest.mark.parametrize(
    ("ruc", "count", "nearest", "expected"),
    [
        ("Urban: Nearer to a major town or city", 5, 300, rules.URBAN),
        ("Smaller rural: Further from a major town or city", 60, 20, rules.URBAN),
        ("Larger rural: Nearer to a major town or city", 3, 400, rules.RURAL),
        ("Larger rural: Nearer to a major town or city", 15, 400, rules.BORDERLINE),
        ("Smaller rural: Nearer to a major town or city", 2, 40, rules.BORDERLINE),
        (None, 2, 400, None),
    ],
)
def test_setting(ruc, count, nearest, expected):
    result, reason = rules.setting(ruc, count, nearest, CFG)
    assert result == expected
    assert reason


def test_noise_falls_6db_per_doubling():
    assert rules.predicted_level(70, 7, CFG) == 70
    assert rules.predicted_level(70, 14, CFG) == pytest.approx(64.0, abs=0.1)
    assert rules.predicted_level(70, 70, CFG) == pytest.approx(50.0, abs=0.1)
    assert rules.predicted_level(70, 2, CFG) == 70  # never louder than at 7 m


def test_impact_bands():
    assert rules.impact(12, CFG) == ("Significant adverse impact likely", "Red")
    assert rules.impact(6, CFG) == ("Adverse impact likely", "Amber")
    assert rules.impact(-3, CFG)[1] == "Green"


def test_assumed_level_from_kva():
    assert rules.assumed_level(60, CFG) == 65
    assert rules.assumed_level(800, CFG) == 75
    assert rules.assumed_level(None, CFG) is None


# ---- reference data prep and full screening -------------------------------------------
@pytest.fixture(scope="module")
def reference(tmp_path_factory):
    import geopandas as gpd
    from shapely.geometry import box

    root = tmp_path_factory.mktemp("ref")
    raw = root / "raw"
    raw.mkdir()

    # ONS output areas
    gpd.GeoDataFrame(
        {"OA21CD": ["E00000001", "E00000002"],
         "RUC21NM": ["Urban: Nearer to a major town or city",
                     "Smaller rural: Further from a major town or city"]},
        geometry=[box(*TOWN), box(*COUNTRY)], crs=27700,
    ).to_file(raw / "oa.gpkg")

    # Addresses: a dense grid in the town, three farms in the countryside
    gx, gy = np.meshgrid(np.arange(TOWN[0] + 50, TOWN[2], 40),
                         np.arange(TOWN[1] + 50, TOWN[3], 40))
    town = pd.DataFrame({"UPRN": 1, "X_COORDINATE": gx.ravel(), "Y_COORDINATE": gy.ravel(),
                         "LATITUDE": 0, "LONGITUDE": 0})
    farms = pd.DataFrame({"UPRN": 2, "X_COORDINATE": [353_400, 353_000, 352_700],
                          "Y_COORDINATE": [451_000, 450_400, 451_600],
                          "LATITUDE": 0, "LONGITUDE": 0})
    csv = pd.concat([town, farms]).to_csv(index=False).encode()
    with zipfile.ZipFile(raw / "uprn.zip", "w") as z:
        z.writestr("osopenuprn.csv", csv)

    # Code-Point Open: nested zip, headerless CSVs under Data/CSV/
    with zipfile.ZipFile(raw / "codepo.zip", "w") as z:
        z.writestr("Data/CSV/zz.csv",
                   '"ZZ1 1AA",10,351000,451000,"E92000001","","","","",""\n'
                   '"ZZ9 9ZZ",10,353500,451500,"E92000001","","","","",""\n')

    out = root / "reference"
    assert prepare.prepare_ruc(raw / "oa.gpkg", out, CFG) == 2
    assert prepare.prepare_addresses(raw / "uprn.zip", out, CFG) == len(town) + 3
    assert prepare.prepare_postcodes(raw / "codepo.zip", out, CFG) == 2
    pd.DataFrame({"name": ["Big Solar"], "technology": ["Solar PV"], "mw": [10.0],
                  "status": ["Operational"], "source": ["REPD"],
                  "e": [353_600], "n": [451_100]}).to_csv(out / "generation.csv", index=False)
    return ReferenceData.load(out)


def test_reference_layers_all_load(reference):
    assert all(reference.available.values())


def test_full_screening(reference):
    sites = pd.DataFrame({
        "Site name": ["Town centre", "Open field", "Farm gate", "By postcode", "Nowhere",
                      "Edinburgh"],
        "Location": ["351000, 451000", "353500, 451500", "353010, 450410", "ZZ1 1AA", "",
                     "NT 25 73"],
        "kVA": ["500", "100", "60", "200", "100", "100"],
        "Noise at 7m": ["", "", "65", "", "", ""],
    })
    result = screen(sites, reference, CFG)
    res = result.results.set_index("Site")

    assert res.loc["Town centre", "Setting"] == rules.URBAN
    assert res.loc["Open field", "Setting"] == rules.RURAL
    assert res.loc["Farm gate", "Setting"] == rules.BORDERLINE   # rural, but a house 14 m away
    assert res.loc["By postcode", "Located by"] == "postcode"
    assert res.loc["By postcode", "Setting"] == rules.URBAN

    # Farm gate: 65 dB at 7 m, nearest address ~14 m -> ~59 dB vs 35 dB borderline night
    farm = res.loc["Farm gate"]
    assert farm["Noise figure"] == "from input file"
    assert farm["Predicted at nearest address dB(A)"] == pytest.approx(59, abs=1)
    assert farm["RAG"] == "Red"

    open_field = res.loc["Open field"]
    assert open_field["Generation within 2 km"] == 1
    assert "Big Solar" in open_field["Nearest existing schemes"]

    problems = result.problems.set_index("Site")["Problem"]
    assert problems["Nowhere"] == "No location given"
    assert "Outside the area covered" in problems["Edinburgh"]


def test_screening_without_reference_data_is_explicit(tmp_path):
    empty = ReferenceData.load(tmp_path)
    sites = pd.DataFrame({"Site": ["A"], "Grid ref": ["SJ 838 980"], "kVA": ["100"]})
    res = screen(sites, empty, CFG).results.iloc[0]
    assert res["Setting"] is None
    assert "not loaded" in res["Why this setting"]
    assert "not loaded" in res["Notes"]


def test_example_file_screens(reference):
    path = Path(__file__).resolve().parents[1] / "app" / "examples" / "example_generator_sites.csv"
    df = read_table(path.read_bytes(), path.name)
    result = screen(df, reference, CFG)
    assert len(result.results) + len(result.problems) == len(df)


def test_site_outside_every_output_area_is_not_assessed(reference):
    """Mixed batch: one site inside the ONS areas, one outside them (classification NaN)."""
    sites = pd.DataFrame({"Site": ["Inside", "Outside"],
                          "Location": ["353500, 451500", "380000, 480000"]})
    res = screen(sites, reference, CFG).results.set_index("Site")
    assert res.loc["Inside", "Setting"] == rules.RURAL
    assert pd.isna(res.loc["Outside", "Setting"])
    assert "not loaded" in res.loc["Outside", "Why this setting"]


def test_excel_report(reference):
    from openpyxl import load_workbook

    from gbeg.screening.report import to_excel

    sites = pd.DataFrame({"Site": ["Farm gate", "Nowhere"],
                          "Location": ["353010, 450410", ""], "kVA": ["60", "1"]})
    result = screen(sites, reference, CFG)
    wb = load_workbook(io.BytesIO(to_excel(result, CFG, reference.available)))
    assert wb.sheetnames == ["Results", "Not screened", "Assumptions"]
    results = wb["Results"]
    header = [c.value for c in results[1]]
    assert results.cell(2, header.index("RAG") + 1).fill.start_color.rgb.endswith("F4C7C3")
    assert wb["Not screened"].cell(2, 2).value == "Nowhere"
    items = {r[0].value: r[1].value for r in wb["Assumptions"].iter_rows(min_row=2)}
    assert "BS 4142" in items["Status"]


def test_location_and_postcode_columns_only():
    """A Location column holding names or addresses falls back to the postcode."""
    from gbeg.screening.inputs import locate_row

    lookup = {"BL11AP": (371951, 409081), "WN11AA": (358016, 405780)}.get
    df = pd.DataFrame({"Location": ["Bolton Depot", "SJ 838 980", "12 High St, Wigan WN1 1AA",
                                    "Nowhere Farm"],
                       "Postcode": ["BL1 1AP", "M1 1AA", "", "ZZ9 9ZZ"]})
    cols = detect_columns(df)
    assert cols.get("location") == "Location" and cols.get("postcode") == "Postcode"
    got = [locate_row(r, cols, lookup) for _, r in df.iterrows()]
    assert [g.method for g in got[:3]] == ["postcode", "grid reference", "postcode"]
    assert got[0].easting == 371951 and got[2].easting == 358016
    assert "Nowhere Farm" in got[3].issue and "ZZ9 9ZZ not found" in got[3].issue


def test_add_sites_fills_file_columns_and_adds_missing_ones():
    from gbeg.screening.inputs import add_sites
    sites = pd.DataFrame({"Name": ["A"], "Easting": ["383800"], "Northing": ["398000"],
                          "kVA": ["500"]})
    cols = detect_columns(sites)
    extra = pd.DataFrame({"Site name": ["Typed", None, ""], "Location": ["SJ 838 980", None, ""],
                          "Generator rating (kVA)": ["250", None, ""]})
    combined, new_cols = add_sites(sites, cols, extra)
    assert len(combined) == 2                      # blank rows are dropped
    assert combined.loc[1, "Name"] == "Typed"      # goes into the file's name column
    assert combined.loc[1, "kVA"] == "250"
    assert new_cols.get("location") == "Location"  # file had none, so one is added
    assert pd.isna(combined.loc[0, "Location"])


def test_typed_sites_screen_with_file(reference):
    from gbeg.screening.inputs import add_sites
    sites = pd.DataFrame({"Site": ["File site"], "Easting": ["351000"], "Northing": ["451000"]})
    extra = pd.DataFrame({"Site name": ["Typed site"], "Location": ["353000, 451000"]})
    combined, cols = add_sites(sites, detect_columns(sites), extra)
    result = screen(combined, reference, CFG, cols)
    assert result.results["Site"].tolist() == ["File site", "Typed site"]
    assert result.problems.empty


# ---- SPEN substation lookup ------------------------------------------------------------
def _spen_reference(reference, tmp_path):
    """The shared reference data plus a small export laid out like the SPEN portal's
    capacity heatmap CSV: one row per transformer."""
    import shutil
    folder = tmp_path / "ref"
    shutil.copytree(reference.folder, folder)
    from gbeg.transform.coords import from_en
    town, field = from_en(351_000, 451_000), from_en(353_500, 451_500)
    town2 = from_en(351_600, 451_600)
    portal = tmp_path / "heatmap.csv"
    pd.DataFrame({
        "name": ["HOOLE RD T1 / HOOLE RD T2"] * 2 + ["FARMSIDE T1", "CREWE T1", "CREWE GT1",
                                                    "EDINBURGH T1"],
        "type": ["Primary", "Primary", "Primary", "Primary", "BSP", "Primary"],
        "voltage": ["[33,11]"] * 4 + ["[132,33]", "[33,11]"],
        "site name (individual site)": ["HOOLE RD T1", "HOOLE RD T2", "FARMSIDE T1",
                                        "CREWE T1", "CREWE GT1", "EDINBURGH T1"],
        "latitude (individual site)": [town.latitude, town.latitude + 0.0001, field.latitude,
                                       town2.latitude, field.latitude, 55.95],
        "longitude (individual site)": [town.longitude, town.longitude, field.longitude,
                                        town2.longitude, field.longitude, -3.19],
    }).to_csv(portal, index=False, encoding="utf-8-sig")
    # Hoole Rd's two transformers merge; Crewe primary and BSP stay apart; Edinburgh clipped
    assert prepare.prepare_spen_sites([portal], folder, CFG) == 4
    return ReferenceData.load(folder)


def test_site_key_ignores_substation_words():
    from gbeg.screening.layers import site_key
    assert site_key("HOOLE ROAD 33/11kV Primary S/S") == "hoole road"
    assert site_key("Hoole Rd substation T1") == "hoole road"
    assert site_key("BURLINGTON ST GT2") == site_key("Burlington Street")


def test_spen_site_found_by_name(reference, tmp_path):
    ref = _spen_reference(reference, tmp_path)
    assert ref.spen_sites_loaded
    assert set(ref.spen_sites["label"]) == {"Hoole Rd", "Farmside", "Crewe (Primary)",
                                            "Crewe (BSP)"}
    assert ref.find_site("Hoole Road Substation")[2] == "Hoole Rd"
    assert ref.find_site("Farmsde")[2] == "Farmside"                 # close spelling
    assert ref.find_site("Crewe") is None                            # two places
    assert ref.find_site("Crewe primary")[2] == "Crewe (Primary)"
    assert ref.find_site("Crewe (BSP)")[2] == "Crewe (BSP)"
    assert ref.find_site("Somewhere else") is None

    sites = pd.DataFrame({"Site name": ["Hoole Road primary", "Farmside", "Unknown", "Given",
                                        "Crewe"],
                          "Location": ["", "", "", "353500, 451500", ""]})
    result = screen(sites, ref, CFG)
    res = result.results.set_index("Site")
    assert res.loc["Hoole Road primary", "Located by"] == "SPEN substation (Hoole Rd)"
    assert res.loc["Hoole Road primary", "Setting"] == rules.URBAN
    assert res.loc["Farmside", "Setting"] == rules.RURAL
    assert res.loc["Given", "Located by"] == "easting/northing"   # a location still wins
    problems = result.problems.set_index("Site")["Problem"]
    assert "not an SPEN substation" in problems["Unknown"]
    assert "more than one SPEN substation (Crewe (BSP), Crewe (Primary))" in problems["Crewe"]
