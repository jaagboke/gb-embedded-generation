"""Write screening results to an Excel workbook a user can attach to their own assessment."""

from __future__ import annotations

import io
from datetime import UTC, datetime

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from gbeg.screening.screen import ScreeningResult

RAG_FILL = {"Red": "F4C7C3", "Amber": "FCE8B2", "Green": "D9EAD3"}
FONT = "Arial"


def _assumptions(cfg: dict, available: dict[str, bool]) -> pd.DataFrame:
    s, n = cfg["setting"], cfg["noise"]
    rows = [
        ("Generated", datetime.now(UTC).date().isoformat()),
        ("Status", "Screening only. Not a BS 4142 assessment or a substitute for one."),
        ("Urban if", f"ONS classes the area as urban, or at least "
                     f"{s['urban_if']['addresses_at_least']} addresses within {s['radius_m']} m"),
        ("Rural if", f"ONS classes the area as rural, fewer than "
                     f"{s['rural_if']['addresses_below']} addresses within {s['radius_m']} m, "
                     f"and the nearest at least {s['rural_if']['nearest_address_at_least_m']} m "
                     "away"),
        ("Borderline", "Anything else"),
        ("Noise prediction", f"L = L{n['reference_distance_m']} - 20 log10(d / "
                             f"{n['reference_distance_m']}); point source, no screening"),
        ("Noise if not given", "; ".join(f"up to {k:,} kVA: {v} dB(A) at 7 m"
                                         for k, v in n["assumed_level_by_kva"])),
        ("Assumed night background", "; ".join(f"{k}: {v} dB LA90" for k, v in
                                               n["assumed_night_background_db"].items())),
        ("Day background", f"night value + {n['day_uplift_db']} dB"),
        ("Impact bands", "; ".join(f"excess ≥ {t} dB: {lbl}" for t, lbl in n["impact_bands"]
                                   if t > -100)),
        ("Nearby generation radius", f"{cfg['nearby_generation']['radius_m']:,} m"),
    ]
    rows += [(f"Data: {k}", "loaded" if v else "NOT LOADED") for k, v in available.items()]
    rows += [
        ("Sources", "ONS Rural Urban Classification 2021; OS Open UPRN; OS Code-Point Open; "
                    "DESNZ REPD; ENWL Embedded Capacity Register. Contains OS data © Crown "
                    "copyright and database right. Contains public sector information licensed "
                    "under the Open Government Licence v3.0."),
    ]
    return pd.DataFrame(rows, columns=["Item", "Value"])


def _format(ws, widths: dict[str, int] | None = None, wrap: set[str] | None = None) -> None:
    header = [c.value for c in ws[1]]
    for cell in ws[1]:
        cell.font = Font(name=FONT, bold=True)
        cell.fill = PatternFill("solid", start_color="E6ECEA")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT)
            if wrap and header[cell.column - 1] in wrap:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    for j, name in enumerate(header, start=1):
        width = (widths or {}).get(name) or min(max(len(str(name)), 10) + 2, 40)
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.freeze_panes = "B2"


def to_excel(result: ScreeningResult, cfg: dict, available: dict[str, bool]) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        result.results.to_excel(xw, sheet_name="Results", index=False)
        problems = result.problems if len(result.problems) else pd.DataFrame(
            {"Row": [], "Site": [], "Location given": [], "Problem": []})
        problems.to_excel(xw, sheet_name="Not screened", index=False)
        _assumptions(cfg, available).to_excel(xw, sheet_name="Assumptions", index=False)

        ws = xw.sheets["Results"]
        _format(ws, {"Site": 32, "Why this setting": 48, "Nearest existing schemes": 60,
                     "Notes": 40, "ONS rural/urban class": 34, "Indicative noise impact": 30},
                wrap={"Why this setting", "Nearest existing schemes", "Notes"})
        header = [c.value for c in ws[1]]
        rag_col = header.index("RAG") + 1
        for row in ws.iter_rows(min_row=2, min_col=rag_col, max_col=rag_col):
            colour = RAG_FILL.get(row[0].value)
            if colour:
                row[0].fill = PatternFill("solid", start_color=colour)
        _format(xw.sheets["Not screened"], {"Site": 32, "Location given": 30, "Problem": 70})
        _format(xw.sheets["Assumptions"], {"Item": 28, "Value": 110}, wrap={"Value"})
    return buf.getvalue()
