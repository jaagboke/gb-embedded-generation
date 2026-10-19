"""Draw a stratified sample of ECR-REPD candidate pairs for hand labelling.

Each ECR record's best candidate is put in a stratum by the rule it got (matched_name,
matched_location, weak_candidate, capacity_mismatch, nearby_other_technology). A fixed
number is drawn from each stratum with a fixed seed, so the sample is reproducible and
every rule, accepted or rejected, gets measured. The stratum sizes are stored with the
sample so precision can be weighted back to the whole population (gbeg.evaluate).

The labeller sees both records side by side, the distance and map links, but not the
rule or the name-similarity score, so the labels are not anchored on the matcher's view.
"""

from __future__ import annotations

import argparse
from pathlib import Path

import duckdb
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from gbeg import pipeline

SAMPLE_SIZES = {
    "matched_name": 15,
    "matched_location": 20,
    "weak_candidate": 10,
    "capacity_mismatch": 10,
    "nearby_other_technology": 5,
}
SEED = 20261013
LABELS = ("same", "different", "unsure")

VISIBLE = [
    ("pair_id", "Pair", 7),
    ("ecr_site", "ENWL site", 30),
    ("ecr_operator", "ENWL operator", 24),
    ("ecr_technology", "ENWL technology (as published)", 34),
    ("ecr_mw", "ENWL MW", 9),
    ("ecr_status", "ENWL status", 20),
    ("ecr_postcode", "ENWL postcode", 11),
    ("repd_site", "REPD site", 38),
    ("repd_operator", "REPD operator", 24),
    ("repd_technology", "REPD technology", 20),
    ("repd_mw", "REPD MW", 9),
    ("repd_status", "REPD status", 20),
    ("repd_postcode", "REPD postcode", 11),
    ("distance_m", "Distance (m)", 11),
    ("ecr_map", "ENWL map", 10),
    ("repd_map", "REPD map", 10),
    ("label", "Label", 11),
    ("notes", "Notes", 30),
]
HIDDEN = [
    "ecr_id", "repd_id", "stratum", "stratum_size", "name_similarity",
    "capacity_ratio", "family_match",
]

SAMPLE_SQL = """
    SELECT
        m.ecr_id, m.repd_id, m.match_rule AS stratum,
        count(*) OVER (PARTITION BY m.match_rule) AS stratum_size,
        m.distance_m, m.name_similarity, m.capacity_ratio, m.family_match,
        e.site_name AS ecr_site, e.operator AS ecr_operator,
        e.technology_raw AS ecr_technology, e.capacity_mw AS ecr_mw,
        e.lifecycle_raw AS ecr_status, e.postcode AS ecr_postcode,
        e.latitude AS ecr_lat, e.longitude AS ecr_lon,
        r.site_name AS repd_site, r.operator AS repd_operator,
        r.technology_raw AS repd_technology, r.capacity_mw AS repd_mw,
        r.lifecycle_raw AS repd_status, r.postcode AS repd_postcode,
        r.latitude AS repd_lat, r.longitude AS repd_lon
    FROM int_site_matches m
    JOIN int_site_records e ON e.source_record_id = m.ecr_id AND e.source LIKE 'ecr_%'
    JOIN int_site_records r ON r.source_record_id = m.repd_id AND r.source = 'repd'
    WHERE e.technology_family <> 'fossil'
"""


def draw_sample(con: duckdb.DuckDBPyConnection, sizes: dict[str, int] = SAMPLE_SIZES,
                seed: int = SEED) -> pd.DataFrame:
    population = con.execute(SAMPLE_SQL).df().sort_values(["ecr_id", "repd_id"])
    parts = []
    for stratum, n in sizes.items():
        pool = population[population["stratum"] == stratum]
        parts.append(pool.sample(n=min(n, len(pool)), random_state=seed))
    sample = pd.concat(parts).sample(frac=1, random_state=seed).reset_index(drop=True)
    sample.insert(0, "pair_id", range(1, len(sample) + 1))
    return sample


def _map_url(lat: float, lon: float) -> str | None:
    if pd.isna(lat) or pd.isna(lon):
        return None
    return f"https://www.google.com/maps/search/?api=1&query={lat:.5f},{lon:.5f}"


def write_workbook(sample: pd.DataFrame, path: Path) -> None:
    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    link = Font(name="Arial", size=10, color="0563C1", underline="single")
    input_fill = PatternFill("solid", start_color="FFFF00")
    header_fill = PatternFill("solid", start_color="D9E1F2")

    wb = Workbook()
    guide = wb.active
    guide.title = "Instructions"
    ws = wb.create_sheet("Pairs")

    rows = [
        ("Match labelling", bold),
        (("Each row on the Pairs sheet is an ENWL ECR record next to the REPD record the "
         "matcher thinks is closest to it. Decide whether they describe the same "
         "physical site."), font),
        ("", font),
        (("What to edit: only the yellow Label column (and Notes if useful). Leave every "
         "other column, including the hidden ones, unchanged."), font),
        (("same = the same site or the same scheme (a phase or extension of it counts as "
         "same)."), font),
        ("different = a different site that happens to be nearby.", font),
        ("unsure = not enough information to tell; add a note saying why.", font),
        (("Tips: compare names, operators and postcodes, then use the two map links. "
         "Distances of 1 to 2 km can still be the same wind or solar farm, because one "
         "source may record the connection point and the other the site centre."), font),
        ("", font),
        (("Example of a filled row: Pair 0 | CLIFTON MARSH | ... | Clifton Marsh Farm, "
         "Preston New Road - Solar Farm | ... | 65 | Label: same | Notes: same name and "
         "capacity"), font),
        ("", font),
    ]
    for i, (text, f) in enumerate(rows, start=1):
        cell = guide.cell(row=i, column=1, value=text)
        cell.font = f
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    n = len(sample)
    label_col = get_column_letter([c for c, _, _ in VISIBLE].index("label") + 1)
    progress = len(rows) + 1
    guide.cell(row=progress, column=1, value="Pairs labelled so far").font = bold
    guide.cell(row=progress, column=2,
               value=f'=COUNTIF(Pairs!{label_col}2:{label_col}{n + 1},"?*")').font = font
    guide.cell(row=progress + 1, column=1, value="Pairs in sample").font = bold
    guide.cell(row=progress + 1, column=2, value=n).font = font
    guide.column_dimensions["A"].width = 110
    guide.column_dimensions["B"].width = 12

    columns = [c for c, _, _ in VISIBLE] + HIDDEN
    headers = [h for _, h, _ in VISIBLE] + HIDDEN
    for j, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=j, value=header)
        cell.font = bold
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    data = sample.copy()
    data["ecr_map"] = [_map_url(a, b) for a, b in zip(data.ecr_lat, data.ecr_lon, strict=True)]
    data["repd_map"] = [_map_url(a, b) for a, b in zip(data.repd_lat, data.repd_lon,
                                                        strict=True)]
    data["label"] = None
    data["notes"] = None

    for i, record in enumerate(data[columns].itertuples(index=False), start=2):
        for j, (col, value) in enumerate(zip(columns, record, strict=True), start=1):
            value = None if (not isinstance(value, str) and pd.isna(value)) else value
            if hasattr(value, "item"):
                value = value.item()
            cell = ws.cell(row=i, column=j)
            if col in ("ecr_map", "repd_map") and value:
                cell.value = "open map"
                cell.hyperlink = value
                cell.font = link
            else:
                cell.value = value
                cell.font = font
            if col in ("label", "notes"):
                cell.fill = input_fill

    for j, (_, _, width) in enumerate(VISIBLE, start=1):
        ws.column_dimensions[get_column_letter(j)].width = width
    for j in range(len(VISIBLE) + 1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(j)].hidden = True

    validation = DataValidation(type="list", formula1=f'"{",".join(LABELS)}"',
                                allow_blank=True)
    validation.add(f"{label_col}2:{label_col}{n + 1}")
    ws.add_data_validation(validation)
    ws[f"{label_col}1"].comment = Comment("Choose same, different or unsure.", "gbeg")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(VISIBLE))}{n + 1}"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Write a labelling sample of match pairs")
    parser.add_argument("--db", type=Path, default=Path("data/processed/gbeg.duckdb"))
    parser.add_argument("--out", type=Path, default=Path("labels/match_sample.xlsx"))
    args = parser.parse_args()
    if args.out.exists():
        raise SystemExit(f"{args.out} already exists; it may hold labels. Move it first.")
    with duckdb.connect(str(args.db)) as con:
        pipeline.build(con)
        sample = draw_sample(con)
    write_workbook(sample, args.out)
    print(f"Wrote {len(sample)} pairs to {args.out}")
    print(sample["stratum"].value_counts().to_string())


if __name__ == "__main__":
    main()
